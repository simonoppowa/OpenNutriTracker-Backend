#!/usr/bin/env python3
"""Convert the Bundeslebensmittelschluessel (BLS 4.0) Excel export into
CSV files matching the multi-source Supabase schema
(multisource_food_schema.mermaid) - the same tables fdc_to_ont_csv.py
fills for FDC.

Emitted CSVs:

    food_source.csv       (1 row: bls)
    bls_component.csv     (full BLS component catalog)
    nutrient_mapping.csv  (source='bls' rows only)
    food.csv              (ids start at 10_000_000; English names)
    food_translation.csv  (native German names, locale 'de')
    food_nutrient.csv     (canonical nutrients, units converted)

NOT emitted: nutrient.csv - the canonical nutrient list is shared and must
already be loaded (it ships with the FDC conversion). Also no
food_portion/food_image: BLS 4.0 contains neither portions nor photos.

Usage:
    # with the xlsx files auto-found under ./source/bls (download_bls.py):
    python bls_to_ont_csv.py --output <out_dir>

    # or point at them explicitly:
    python bls_to_ont_csv.py \
        --data BLS_4_0_Daten_2025_DE.xlsx \
        --components BLS_4_0_Components_DE_EN.xlsx \
        --output <out_dir>

Typical pipeline:
    python download_bls.py --output source        # -> ./source/bls
    python bls_to_ont_csv.py --output ../../out/bls  # reads ./source/bls

Requires: pip install openpyxl

Import after (or alongside) the FDC CSVs; append, don't truncate:
    \\copy food_source from 'food_source.csv' csv header
    \\copy bls_component from 'bls_component.csv' csv header
    \\copy nutrient_mapping from 'nutrient_mapping.csv' csv header
    \\copy food from 'food.csv' csv header
    \\copy food_translation (food_id, locale, description, source) from 'food_translation.csv' csv header
    \\copy food_nutrient from 'food_nutrient.csv' csv header
    refresh materialized view food_summary;

Notes
-----
* food.id starts at FOOD_ID_BASE (10_000_000) so BLS ids can never
  collide with FDC ids (which reuse fdc_id, all < 10_000_000).
* BLS has no trans-fat component -> trans_fat_100 stays NULL for BLS
  foods (correctly maps to `double?` in the app).
* Vitamin B6 is reported in ug by BLS but canonical unit is mg ->
  unit_factor 0.001. Everything else already matches.
* BLS 4.0 is Open Data under CC BY 4.0 (Max Rubner-Institut) - free to
  use and redistribute with attribution (cite MRI 2025, DOI
  10.25826/Data20251217-134202-0). Download with download_bls.py.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("error: openpyxl missing - run: pip install openpyxl")

FOOD_ID_BASE = 10_000_000          # keep clear of FDC fdc_ids
NUTRIENT_ID_BASE = 20_000_000
DEFAULT_SOURCE_DIR = Path("source") / "bls"  # what download_bls.py fills

# ---------------------------------------------------------------------------
# Canonical nutrient ids (must match fdc_to_ont_csv.py) mapped to
# BLS component codes, in priority order (first present per food wins).
# unit_factor converts the BLS unit to the canonical unit.
# ---------------------------------------------------------------------------
BLS_MAPPING = [
    # (canonical id, [(bls_code, unit_factor), ...])
    (1,  [("ENERCC", 1)]),            # kcal
    (2,  [("CHO", 1)]),               # available carbohydrate, g
    (3,  [("FAT", 1)]),               # g
    (4,  [("PROT625", 1)]),           # protein Nx6.25, g
    (5,  [("SUGAR", 1)]),             # total sugars, g
    (6,  [("FASAT", 1)]),             # g
    (7,  [("FIBT", 1)]),              # total fibre, g
    (8,  [("FAMS", 1)]),              # g
    (9,  [("FAPU", 1)]),              # g
    # 10: trans fat - no BLS component, stays NULL
    (11, [("CHORL", 1)]),             # cholesterol, mg
    (12, [("NA", 1)]),                # mg
    (13, [("K", 1)]),                 # mg
    (14, [("MG", 1)]),                # mg
    (15, [("CA", 1)]),                # mg
    (16, [("FE", 1)]),                # mg
    (17, [("ZN", 1)]),                # mg
    (18, [("P", 1)]),                 # mg
    (19, [("VITAA", 1), ("VITA", 1)]),  # ug RAE, fallback ug RE
    (20, [("VITC", 1)]),              # mg
    (21, [("VITD", 1)]),              # ug
    (22, [("VITB6", 0.001)]),         # BLS: ug -> canonical mg
    (23, [("VITB12", 1)]),            # ug
    (24, [("NIA", 1)]),               # mg
]

ORIGIN_BY_HERKUNFT = {
    "analyse": "analysis",
    "literatur": "literature",
    "nährstoffdatenbank": "literature",   # taken from another food database
    "labelangabe": "label",
    "formelberechnung": "calculated",
    "aggregation": "calculated",
    "reskalierung": "calculated",
    "rezeptberechnung": "calculated",
    "logische annahme": "calculated",
    "logische null": "calculated",
    "-": "",
}


def short_title(description: str) -> str:
    """Concise display name: text before the first comma, capped at 80."""
    head = description.split(",")[0].strip()
    return (head or description.strip())[:80]


def parse_amount(value) -> float | None:
    """BLS cells hold floats, dot- or comma-decimal strings, or '-'."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", ".")
    if not s or s in ("-", "n.a."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def convert(data_xlsx: Path, components_xlsx: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    # ---- component catalog -> bls_component.csv ---------------------------
    wb = openpyxl.load_workbook(components_xlsx, read_only=True)
    ws = wb[wb.sheetnames[0]]
    components = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        components.append({
            "code": str(row[1]).strip(),
            "name_de": str(row[2] or "").strip(),
            "name_en": str(row[3] or "").strip(),
            "unit": str(row[4] or "").strip(),
            "component_group": str(row[6] or "").strip(),
            "formula": str(row[7] or "").strip(),
        })
    wb.close()

    # ---- data sheet ---------------------------------------------------------
    wb = openpyxl.load_workbook(data_xlsx, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    # Column layout: 'CODE <label> [unit/100g]', 'CODE Datenherkunft',
    # 'CODE Referenz'. Locate the value/herkunft/referenz columns per code.
    value_col, herkunft_col, referenz_col = {}, {}, {}
    for i, h in enumerate(header[3:], start=3):
        if not h:
            continue
        h = str(h)
        code = h.split(" ", 1)[0]
        if h.endswith("Datenherkunft"):
            herkunft_col[code] = i
        elif h.endswith("Referenz"):
            referenz_col[code] = i
        elif re.search(r"\[.+/100g\]$", h):
            value_col[code] = i

    # canonical id -> [(code, factor, priority), ...] for codes in the file
    lookup = {}
    missing = []
    for nid, options in BLS_MAPPING:
        found = [(c, f, p) for p, (c, f) in enumerate(options) if c in value_col]
        if found:
            lookup[nid] = found
        else:
            missing.append((nid, options[0][0]))
    for nid, code in missing:
        print(f"  note: no BLS column for canonical nutrient {nid} ({code})",
              file=sys.stderr)

    foods, translations, food_nutrients = [], [], []
    seen_codes: set[str] = set()
    unknown_origins: set[str] = set()
    next_fn_id = NUTRIENT_ID_BASE

    for row in rows:
        bls_code = str(row[0] or "").strip()
        name_de = str(row[1] or "").strip()
        name_en = str(row[2] or "").strip()
        if not bls_code or bls_code in seen_codes:
            continue
        seen_codes.add(bls_code)
        food_id = FOOD_ID_BASE + len(foods)
        foods.append({
            "id": food_id,
            "source": "bls",
            "source_code": bls_code,
            "description": name_en or name_de,   # English name; DE fallback
            "short_title": short_title(name_en or name_de),
            "food_category_id": "",
            "publication_date": "",
        })
        if name_de:
            translations.append({
                "food_id": food_id,
                "locale": "de",
                "description": name_de,
                "source": "native",
            })
        for nid, options in lookup.items():
            for code, factor, _prio in options:
                amount = parse_amount(row[value_col[code]])
                if amount is None:
                    continue
                herkunft = str(row[herkunft_col.get(code, 0)] or "").strip() \
                    if code in herkunft_col else ""
                key = herkunft.lower()
                if key in ORIGIN_BY_HERKUNFT:
                    origin = ORIGIN_BY_HERKUNFT[key]
                elif herkunft:
                    unknown_origins.add(herkunft)
                    origin = "calculated"
                else:
                    origin = ""
                ref = str(row[referenz_col.get(code, 0)] or "").strip() \
                    if code in referenz_col else ""
                food_nutrients.append({
                    "id": next_fn_id,
                    "food_id": food_id,
                    "nutrient_id": nid,
                    "amount": round(amount * factor, 6),
                    "origin": origin,
                    "reference": "" if ref == "-" else ref,
                    "data_points": "",
                    "min": "",
                    "max": "",
                })
                next_fn_id += 1
                break                      # first available option wins
    wb.close()

    if unknown_origins:
        print(f"  note: unmapped Datenherkunft values -> 'calculated': "
              f"{sorted(unknown_origins)[:8]}", file=sys.stderr)

    # ------------------------------------------------------------------ write
    def write(name: str, rows_, columns: list[str]) -> None:
        with open(out_dir / f"{name}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows_)
        print(f"  {name:20s} {len(rows_):>7,} rows")

    print("Writing Supabase CSVs:")
    write("food_source",
          [{"code": "bls",
            "name": "Bundeslebensmittelschluessel 4.0 (Max Rubner-Institut)",
            "license_note": "CC BY 4.0 - attribute Max Rubner-Institut (2025), "
                            "DOI 10.25826/Data20251217-134202-0"}],
          ["code", "name", "license_note"])
    write("bls_component", components,
          ["code", "name_de", "name_en", "unit", "component_group", "formula"])
    write("nutrient_mapping",
          [{"source": "bls", "source_component_code": c,
            "nutrient_id": nid, "unit_factor": f}
           for nid, opts in BLS_MAPPING for c, f in opts],
          ["source", "source_component_code", "nutrient_id", "unit_factor"])
    write("food", foods,
          ["id", "source", "source_code", "description", "short_title",
           "food_category_id", "publication_date"])
    write("food_translation", translations,
          ["food_id", "locale", "description", "source"])
    write("food_nutrient", food_nutrients,
          ["id", "food_id", "nutrient_id", "amount", "origin",
           "reference", "data_points", "min", "max"])

    n_energy = len({r["food_id"] for r in food_nutrients if r["nutrient_id"] == 1})
    print(f"\nDone. {len(foods):,} BLS foods ({n_energy:,} with energy) -> {out_dir}")


def discover(source_dir: Path, pattern: str, label: str) -> Path:
    """Find one xlsx under source_dir matching pattern (download_bls.py
    extracts into ./source/bls)."""
    hits = sorted(source_dir.rglob(pattern))
    if not hits:
        sys.exit(f"error: no {label} ({pattern}) under {source_dir} - "
                 f"run download_bls.py first, or pass --data/--components")
    return hits[0]


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Convert BLS 4.0 xlsx to multi-source Supabase CSVs")
    ap.add_argument("--data", type=Path, default=None,
                    help="BLS_4_0_Daten_*.xlsx (default: found under "
                         "--source-dir)")
    ap.add_argument("--components", type=Path, default=None,
                    help="BLS_4_0_Components_DE_EN.xlsx (default: found under "
                         "--source-dir)")
    ap.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR,
                    help="folder from download_bls.py holding the xlsx files "
                         "(default: ./source/bls)")
    ap.add_argument("--output", required=True, type=Path)
    args = ap.parse_args()

    data = args.data or discover(args.source_dir, "*Daten*.xlsx", "BLS data file")
    components = args.components or discover(
        args.source_dir, "*Components*.xlsx", "BLS components file")
    for f in (data, components):
        if not f.exists():
            sys.exit(f"error: {f} not found")
    print(f"data:       {data}")
    print(f"components: {components}")
    convert(data, components, args.output)


if __name__ == "__main__":
    main()
