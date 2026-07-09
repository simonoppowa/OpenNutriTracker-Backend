#!/usr/bin/env python3
"""Convert the Anuvaad Indian Nutrient Databank (INDB) Excel export into
CSV files matching the multi-source Supabase schema - same tables the
FDC and BLS converters fill.

Emitted CSVs:

    food_source.csv       (1 row: indb)
    nutrient_mapping.csv  (source='indb' rows, documents the mapping)
    food.csv              (ids start at 20_000_000)
    food_alias.csv        (romanized Indian names from parentheses,
                           e.g. 'Garam Chai' for 'Hot tea (Garam Chai)')
    food_portion.csv      (one serving per food, gram weight derived
                           from per-serving vs per-100g energy ratio)
    food_nutrient.csv     (canonical nutrients, units converted)

NOT emitted: nutrient.csv (shared canonical list, ships with the FDC
conversion) and measure_unit.csv (portions reference FDC's unit 9999
'undetermined'; load the FDC CSVs first).

Usage:
    python indb_to_ont_csv.py --data Anuvaad_INDB_2024.11.xlsx \
                                      --output <out_dir>

Requires: pip install openpyxl

Import (append) after FDC:
    \\copy food_source from 'food_source.csv' csv header
    \\copy nutrient_mapping from 'nutrient_mapping.csv' csv header
    \\copy food from 'food.csv' csv header
    \\copy food_alias (food_id, locale, alias, source) from 'food_alias.csv' csv header
    \\copy food_portion from 'food_portion.csv' csv header
    \\copy food_nutrient from 'food_nutrient.csv' csv header
    refresh materialized view food_summary;

Notes
-----
* id ranges: FDC uses fdc_ids (< 10M), BLS 10M+, INDB starts at 20M
  (food) / 30M (food_nutrient) - no collisions.
* Unit conversions: sfa/mufa/pufa come in mg -> g (x0.001); vitamin D
  is D2 + D3 summed (both ug); everything else already canonical.
* 'freesugar_g' is mapped to canonical 'sugars' - INDB reports free
  sugars, the closest available concept.
* No trans-fat and no vitamin B12 columns in INDB -> both stay NULL
  (map to nullable `double?` in the app).
* Serving gram weight is derived: unit_serving_energy / energy * 100,
  falling back to the protein / carb / fat ratio when energy is 0.
* INDB (https://anuvaad.org.in) is CC BY 4.0 - attribute in the app.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("error: openpyxl missing - run: pip install openpyxl")

FOOD_ID_BASE = 20_000_000
NUTRIENT_ID_BASE = 30_000_000
UNDETERMINED_UNIT_ID = 9999          # FDC measure_unit 'undetermined'

# canonical nutrient id -> (INDB column name, unit_factor)
# Column indices are resolved from the header row at runtime.
INDB_MAPPING = [
    (1,  "energy_kcal",    1),
    (2,  "carb_g",         1),
    (3,  "fat_g",          1),
    (4,  "protein_g",      1),
    (5,  "freesugar_g",    1),        # free sugars ~ total sugars (closest)
    (6,  "sfa_mg",         0.001),    # mg -> g
    (7,  "fibre_g",        1),
    (8,  "mufa_mg",        0.001),    # mg -> g
    (9,  "pufa_mg",        0.001),    # mg -> g
    # 10 trans fat: not in INDB
    (11, "cholesterol_mg", 1),
    (12, "sodium_mg",      1),
    (13, "potassium_mg",   1),
    (14, "magnesium_mg",   1),
    (15, "calcium_mg",     1),
    (16, "iron_mg",        1),
    (17, "zinc_mg",        1),
    (18, "phosphorus_mg",  1),
    (19, "vita_ug",        1),
    (20, "vitc_mg",        1),
    # 21 vitamin D: special-cased as vitd2_ug + vitd3_ug below
    (22, "vitb6_mg",       1),
    # 23 vitamin B12: not in INDB
    (24, "vitb3_mg",       1),        # niacin
]
VITD_ID = 21

# per-food provenance -> food_nutrient.origin
ORIGIN_BY_PRIMARYSOURCE = {
    "asc_manual": "literature",
    "bfp_manual": "literature",
    "open_source_recipes": "calculated",
}

ALIAS_RE = re.compile(r"\(([^()]{2,60})\)\s*$")


def short_title(description: str) -> str:
    """Concise display name: drop a trailing '(romanized)' and take the
    text before the first comma, capped at 80 chars."""
    base = ALIAS_RE.sub("", description).strip()
    head = base.split(",")[0].strip()
    return (head or base or description.strip())[:80]


def val(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f


def convert(data_xlsx: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(data_xlsx, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(rows)]
    col = {name: i for i, name in enumerate(header)}

    for required in ("food_code", "food_name", "energy_kcal", "servings_unit",
                     "unit_serving_energy_kcal"):
        if required not in col:
            sys.exit(f"error: column '{required}' not found in {data_xlsx}")

    mapping = [(nid, name, f) for nid, name, f in INDB_MAPPING if name in col]

    foods, aliases, portions, food_nutrients = [], [], [], []
    seen: set[str] = set()
    next_fn_id = NUTRIENT_ID_BASE

    for row in rows:
        code = str(row[col["food_code"]] or "").strip()
        name = str(row[col["food_name"]] or "").strip()
        if not code or code in seen or not name:
            continue
        seen.add(code)
        food_id = FOOD_ID_BASE + len(foods)
        origin = ORIGIN_BY_PRIMARYSOURCE.get(
            str(row[col["primarysource"]] or "").strip().lower(), "")

        foods.append({
            "id": food_id,
            "source": "indb",
            "source_code": code,
            "description": name,
            "short_title": short_title(name),
            "food_category_id": "",
            "publication_date": "",
        })

        # Romanized Indian name in trailing parentheses -> searchable alias.
        m = ALIAS_RE.search(name)
        if m:
            aliases.append({
                "food_id": food_id,
                "locale": "hi-Latn",
                "alias": m.group(1).strip(),
                "source": "derived",
            })

        # ---- nutrients (per 100 g) --------------------------------------
        def add_nutrient(nid: int, amount: float) -> None:
            nonlocal next_fn_id
            food_nutrients.append({
                "id": next_fn_id,
                "food_id": food_id,
                "nutrient_id": nid,
                "amount": round(amount, 6),
                "origin": origin,
                "reference": "",
                "data_points": "", "min": "", "max": "",
            })
            next_fn_id += 1

        for nid, colname, factor in mapping:
            amount = val(row[col[colname]])
            if amount is not None and amount >= 0:
                add_nutrient(nid, amount * factor)

        d2 = val(row[col["vitd2_ug"]]) if "vitd2_ug" in col else None
        d3 = val(row[col["vitd3_ug"]]) if "vitd3_ug" in col else None
        if d2 is not None or d3 is not None:
            add_nutrient(VITD_ID, (d2 or 0) + (d3 or 0))

        # ---- serving -> food_portion --------------------------------------
        unit = str(row[col["servings_unit"]] or "").strip()
        kcal100 = val(row[col["energy_kcal"]])
        kcal_serv = val(row[col["unit_serving_energy_kcal"]])
        grams = None
        if kcal100 and kcal_serv:
            grams = kcal_serv / kcal100 * 100
        else:                                   # fallback: macro ratio
            for base_col, serv_col in (("protein_g", "unit_serving_protein_g"),
                                       ("carb_g", "unit_serving_carb_g"),
                                       ("fat_g", "unit_serving_fat_g")):
                b, s = val(row[col[base_col]]), val(row[col[serv_col]])
                if b and s:
                    grams = s / b * 100
                    break
        if unit and grams and grams > 0:
            portions.append({
                "id": 20_000_000 + len(portions),   # clear of FDC portion ids
                "food_id": food_id,
                "amount": 1,
                "measure_unit_id": UNDETERMINED_UNIT_ID,
                "portion_description": f"1 {unit}",
                "modifier": "",
                "seq_num": 1,
                "gram_weight": round(grams, 2),
            })
    wb.close()

    # ------------------------------------------------------------------ write
    def write(name: str, rows_, columns: list[str]) -> None:
        with open(out_dir / f"{name}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows_)
        print(f"  {name:20s} {len(rows_):>7,} rows")

    print("Writing Supabase CSVs:")
    write("food_source",
          [{"code": "indb",
            "name": "Anuvaad Indian Nutrient Databank (INDB)",
            "license_note": "CC BY 4.0 - attribute anuvaad.org.in"}],
          ["code", "name", "license_note"])
    write("nutrient_mapping",
          [{"source": "indb", "source_component_code": c,
            "nutrient_id": nid, "unit_factor": f}
           for nid, c, f in INDB_MAPPING] +
          [{"source": "indb", "source_component_code": "vitd2_ug+vitd3_ug",
            "nutrient_id": VITD_ID, "unit_factor": 1}],
          ["source", "source_component_code", "nutrient_id", "unit_factor"])
    write("food", foods,
          ["id", "source", "source_code", "description", "short_title",
           "food_category_id", "publication_date"])
    write("food_alias", aliases, ["food_id", "locale", "alias", "source"])
    write("food_portion", portions,
          ["id", "food_id", "amount", "measure_unit_id",
           "portion_description", "modifier", "seq_num", "gram_weight"])
    write("food_nutrient", food_nutrients,
          ["id", "food_id", "nutrient_id", "amount", "origin",
           "reference", "data_points", "min", "max"])

    n_energy = len({r["food_id"] for r in food_nutrients if r["nutrient_id"] == 1})
    print(f"\nDone. {len(foods):,} INDB foods ({n_energy:,} with energy, "
          f"{len(portions):,} servings, {len(aliases):,} aliases) -> {out_dir}")


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Convert Anuvaad INDB xlsx to multi-source Supabase CSVs")
    ap.add_argument("--data", required=True, type=Path,
                    help="Anuvaad_INDB_*.xlsx")
    ap.add_argument("--output", required=True, type=Path)
    args = ap.parse_args()
    if not args.data.exists():
        sys.exit(f"error: {args.data} not found")
    convert(args.data, args.output)


if __name__ == "__main__":
    main()
